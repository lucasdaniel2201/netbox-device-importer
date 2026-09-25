"""Leitura, validacao e geracao do modelo de planilha de cameras para o NetBox.

Aqui as colunas sao as do **NetBox**:

- `Site` e `Papel` sao obrigatoriamente do NetBox e viram **lista suspensa** na
  celula, alimentada pela aba de apoio "Listas" (visivel e filtravel, porque a lista
  de sites e grande e o filtro do Excel tem busca).
- `Firmware` e `Servidor (NVR)` vao para os `comments` do Device (os custom
  fields estao bloqueados para o token atual).
- `Papel` e opcional: em branco, vale o papel padrao escolhido na tela. `Status` nao
  vem da planilha.
- O device-type e resolvido por (fabricante, modelo); se faltar, o app cria.

O nome da camera NAO leva tratamento agressivo: o NetBox aceita UTF-8. So
normalizamos espacos, e usamos uma chave sem acento/caixa para comparar.
"""

from __future__ import annotations

import csv
import re
import sys
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

IP_RE = re.compile(r"^\d{1,3}(\.\d{1,3}){3}$")
MAC_RE = re.compile(r"^([0-9A-Fa-f]{2}[:-]){5}[0-9A-Fa-f]{2}$")

# Limite do campo `name` do device, lido do OPTIONS da instancia (NetBox 4.3.6).
MAX_DEVICE_NAME_LENGTH = 64

# Aba de apoio (visivel, com filtro) que guarda as listas usadas pelos dropdowns.
LIST_SHEET_TITLE = "Listas"
# Ate onde o dropdown vale (o usuario pode acrescentar linhas na planilha).
DROPDOWN_ROWS = 500
DATA_SHEET_TITLE = "Cameras"


@dataclass(frozen=True)
class ColumnSpec:
    """Definicao de uma coluna da planilha."""

    key: str
    label: str
    aliases: tuple[str, ...] = ()
    sample: str = ""
    hint: str = ""
    required_value: bool = False
    fixed: bool = False


COLUMN_SPECS: tuple[ColumnSpec, ...] = (
    ColumnSpec(
        key="Name",
        label="Nome do host",
        aliases=("Name", "Nome"),
        sample="AME-STP2-CAM1-RECEPCAOINTERNA-PIRAPORA-SP-LTL",
        hint="Nome do device no NetBox (obrigatorio, ate 64 caracteres).",
        required_value=True,
        fixed=True,
    ),
    ColumnSpec(
        key="Site",
        label="Site",
        aliases=("Site", "Local", "Unidade"),
        sample="AME VILANOVA",
        hint="Site do NetBox (obrigatorio). No modelo sai como lista suspensa.",
        required_value=True,
        fixed=True,
    ),
    ColumnSpec(
        key="Role",
        label="Papel (role)",
        aliases=("Papel", "Role", "Funcao", "Função"),
        sample="Camera",
        hint="Papel do device no NetBox. Em branco = o papel escolhido na tela.",
    ),
    ColumnSpec(
        key="IP",
        label="IP",
        aliases=("IP/Nome", "IP da camera"),
        sample="10.100.20.58",
        hint="IP da camera. Em branco = a camera e documentada sem endereco.",
        fixed=True,
    ),
    ColumnSpec(
        key="Vendor",
        label="Fabricante",
        aliases=("Vendor", "Fabricante:"),
        sample="Hikvision",
        hint="Fabricante: compoe o device-type junto com o modelo.",
    ),
    ColumnSpec(
        key="Model",
        label="Modelo",
        aliases=("Model",),
        sample="DS-2CD3666G2T-IZS",
        hint="Modelo: compoe o device-type junto com o fabricante.",
    ),
    ColumnSpec(
        key="Firmware",
        label="Firmware",
        sample="V5.7.55 build 241108",
        hint="Vai para os comments do device.",
    ),
    ColumnSpec(
        key="MAC address",
        label="Endereço MAC",
        aliases=("MAC address", "MAC"),
        sample="AA-BB-CC-DD-EE-FF",
        hint="MAC da interface principal do device.",
    ),
    ColumnSpec(
        key="Server",
        label="Servidor (NVR)",
        aliases=("Server", "NVR", "Servidor"),
        sample="NVR-1 (10.100.20.4)",
        hint="NVR de origem. Vai para os comments do device.",
    ),
    ColumnSpec(
        key="Description",
        label="Descrição",
        aliases=("Descricao", "Description"),
        sample="Camera da recepcao interna",
        hint="Descricao do device no NetBox.",
    ),
)

COLUMNS_BY_KEY = {spec.key: spec for spec in COLUMN_SPECS}
TARGET_FIELDS = [spec.key for spec in COLUMN_SPECS]
FIXED_COLUMNS = [spec.key for spec in COLUMN_SPECS if spec.fixed]
OPTIONAL_COLUMNS = [spec.key for spec in COLUMN_SPECS if not spec.fixed]
SITE_KEY = "Site"
ROLE_KEY = "Role"


def normalize_for_match(value: str) -> str:
    """Chave de comparacao: sem acento, sem caixa e sem espaco duplicado."""
    decomposed = unicodedata.normalize("NFKD", value or "")
    without_accents = "".join(c for c in decomposed if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", without_accents).strip().casefold()


def normalize_vendor(vendor: str) -> str:
    """Padroniza o fabricante (a planilha as vezes vem em caixa alta)."""
    value = (vendor or "").strip()
    if value.upper() == "HIKVISION":
        return "Hikvision"
    if value.upper() == "DAHUA":
        return "Dahua"
    return value


def normalize_device_name(name: str) -> str:
    """Ajuste leve do nome: so espaco. O NetBox aceita acentos e pontuacao."""
    return re.sub(r"\s+", " ", (name or "").strip())


def site_index(sites: list[str] | None) -> dict[str, str]:
    """Mapa chave normalizada -> nome do site como esta no NetBox."""
    return {normalize_for_match(site): site for site in (sites or []) if site}


def role_index(roles: list[str] | None) -> dict[str, str]:
    """Mapa chave normalizada -> nome do papel (role) como esta no NetBox."""
    return {normalize_for_match(role): role for role in (roles or []) if role}


@dataclass
class RowResult:
    """Uma linha da planilha apos leitura, normalizacao e validacao."""

    row_number: int
    name_original: str
    name_normalized: str
    cells: dict[str, str]
    issues: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)

    @property
    def valid(self) -> bool:
        return not self.issues


@dataclass
class SheetResult:
    filename: str
    rows: list[RowResult]
    columns_found: list[str] = field(default_factory=list)
    file_warnings: list[str] = field(default_factory=list)
    skipped_empty: int = 0

    @property
    def valid_rows(self) -> list[RowResult]:
        return [row for row in self.rows if row.valid]

    @property
    def rows_with_errors(self) -> list[RowResult]:
        return [row for row in self.rows if not row.valid]


def template_columns(selected_optional: list[str] | None = None) -> list[str]:
    """Colunas do modelo: fixas + opcionais escolhidas (ordem canonica)."""
    optional = OPTIONAL_COLUMNS if selected_optional is None else selected_optional
    chosen = set(FIXED_COLUMNS) | set(optional)
    return [key for key in TARGET_FIELDS if key in chosen]


def writable_optional_columns() -> list[ColumnSpec]:
    """Colunas opcionais oferecidas no popup de geracao do modelo."""
    return [COLUMNS_BY_KEY[key] for key in OPTIONAL_COLUMNS]


def _normalize_header(value: str) -> str:
    """Cabecalho sem acentos e sem espacos duplicados."""
    return re.sub(r"\s+", " ", normalize_for_match(value)).strip()


def _alias_candidates(spec: ColumnSpec) -> list[str]:
    """Nomes de coluna aceitos para um campo destino (rotulo, chave e alias)."""
    return [spec.label, spec.key, *spec.aliases]


def _map_headers(header_cells: list[str]) -> tuple[dict[str, int], list[str]]:
    """Mapeia cabecalhos para campos destino. Retorna (indices, colunas ausentes)."""
    wanted: dict[str, list[str]] = {
        spec.key: [_normalize_header(alias) for alias in _alias_candidates(spec)]
        for spec in COLUMN_SPECS
    }
    normalized_headers = [_normalize_header(cell) for cell in header_cells]

    index_by_target: dict[str, int] = {}
    for key, aliases in wanted.items():
        for header_index, header in enumerate(normalized_headers):
            if header in aliases:
                index_by_target[key] = header_index
                break

    missing = [key for key in TARGET_FIELDS if key not in index_by_target]
    return index_by_target, missing


def _row_to_cells(values: list[str], index_by_target: dict[str, int]) -> dict[str, str]:
    cells: dict[str, str] = {}
    for key, idx in index_by_target.items():
        cells[key] = values[idx].strip() if idx < len(values) else ""
    for key in TARGET_FIELDS:
        cells.setdefault(key, "")
    return cells


def _validate_row(
    row_number: int,
    cells: dict[str, str],
    dup_lines: dict[str, list[int]],
    sites: dict[str, str],
    roles: dict[str, str],
) -> RowResult:
    name_original = cells.get("Name", "")
    name_normalized = normalize_device_name(name_original)
    result = RowResult(
        row_number=row_number,
        name_original=name_original,
        name_normalized=name_normalized,
        cells=cells,
    )

    if name_original and name_normalized != name_original:
        result.warnings.append(f"Nome ajustado: '{name_original}' -> '{name_normalized}'")

    name_key = normalize_for_match(name_normalized)
    if not name_normalized:
        result.issues.append("Nome vazio.")
    elif len(name_normalized) > MAX_DEVICE_NAME_LENGTH:
        result.issues.append(
            f"Nome com {len(name_normalized)} caracteres (o NetBox aceita ate "
            f"{MAX_DEVICE_NAME_LENGTH})."
        )
    elif len(dup_lines.get(name_key, [])) > 1:
        outras = [linha for linha in dup_lines[name_key] if linha != row_number]
        result.issues.append(
            "Nome duplicado com a(s) linha(s): "
            + ", ".join(str(linha) for linha in outras)
        )

    site = cells.get(SITE_KEY, "")
    if not site:
        result.issues.append("Site vazio (escolha na lista suspensa do modelo).")
    elif sites and normalize_for_match(site) not in sites:
        result.issues.append(
            f"Site '{site}' nao existe no NetBox (use a lista suspensa do modelo)."
        )

    ip = cells.get("IP", "")
    if ip and (not IP_RE.match(ip) or any(int(part) > 255 for part in ip.split("."))):
        # Em branco e aceito (camera ainda sem endereco); preenchido tem que ser valido.
        result.issues.append(f"IP com formato invalido: '{ip}'.")

    role = (cells.get(ROLE_KEY) or "").strip()
    if role and roles and normalize_for_match(role) not in roles:
        result.issues.append(
            f"Papel '{role}' nao existe no NetBox (use a lista suspensa do modelo)."
        )

    mac = cells.get("MAC address", "")
    if mac and not MAC_RE.match(mac):
        result.warnings.append(f"MAC com formato incomum: '{mac}'.")

    if cells.get("Vendor"):
        cells["Vendor"] = normalize_vendor(cells["Vendor"])

    return result


def _rows_from_table(
    filename: str,
    header_cells: list[str],
    table_rows: list[list[str]],
    file_warnings: list[str],
    sites: list[str] | None,
    roles: list[str] | None,
) -> SheetResult:
    index_by_target, missing = _map_headers(header_cells)

    # Coluna do modelo ausente e erro; o *valor* obrigatorio e outra checagem
    # (coluna "IP" pode vir em branco - a camera fica documentada sem endereco).
    required_missing = [COLUMNS_BY_KEY[key].label for key in FIXED_COLUMNS if key in missing]
    if required_missing:
        raise ValueError(
            "Coluna(s) obrigatoria(s) do modelo nao encontrada(s): "
            + ", ".join(required_missing)
            + ". Confira o cabecalho ou baixe o modelo de exemplo."
        )

    known_sites = site_index(sites)
    known_roles = role_index(roles)
    columns_found = [key for key in TARGET_FIELDS if key in index_by_target]

    skipped_empty = 0
    seen_names: dict[str, list[int]] = {}
    for row_number, values in enumerate(table_rows, start=2):
        cells = _row_to_cells(values, index_by_target)
        if not cells.get("Name", ""):
            skipped_empty += 1
            continue
        key = normalize_for_match(normalize_device_name(cells["Name"]))
        seen_names.setdefault(key, []).append(row_number)

    rows: list[RowResult] = []
    for row_number, values in enumerate(table_rows, start=2):
        cells = _row_to_cells(values, index_by_target)
        if not cells.get("Name", ""):
            continue
        rows.append(_validate_row(row_number, cells, seen_names, known_sites, known_roles))

    return SheetResult(
        filename=filename,
        rows=rows,
        columns_found=columns_found,
        file_warnings=file_warnings,
        skipped_empty=skipped_empty,
    )


def read_xlsx(
    path: Path, sites: list[str] | None = None, roles: list[str] | None = None
) -> SheetResult:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook.active
        raw_rows = list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()

    if not raw_rows:
        raise ValueError("A planilha esta vazia.")

    header_cells = [str(cell).strip() if cell is not None else "" for cell in raw_rows[0]]
    table_rows = [
        [str(cell) if cell is not None else "" for cell in row] for row in raw_rows[1:]
    ]
    return _rows_from_table(path.name, header_cells, table_rows, [], sites, roles)


def read_csv(
    path: Path, sites: list[str] | None = None, roles: list[str] | None = None
) -> SheetResult:
    encodings = ["utf-8-sig", "cp1252", "latin-1"]
    content: list[list[str]] = []
    used_encoding = None
    for encoding in encodings:
        try:
            with path.open("r", encoding=encoding, newline="") as file:
                content = list(csv.reader(file))
            used_encoding = encoding
            break
        except UnicodeDecodeError:
            continue
    if not content:
        raise ValueError("Nao foi possivel ler o CSV (encoding desconhecido).")

    header_cells = [cell.strip() for cell in content[0]]
    table_rows = [[cell for cell in row] for row in content[1:]]

    file_warnings: list[str] = []
    if used_encoding and used_encoding != "utf-8-sig":
        file_warnings.append(f"CSV lido como {used_encoding}.")
    suspensas = [COLUMNS_BY_KEY[key].label for key in (SITE_KEY, ROLE_KEY) if key in header_cells]
    if suspensas:
        file_warnings.append(
            "Em CSV nao ha lista suspensa: confira a mao as colunas " + ", ".join(suspensas) + "."
        )

    return _rows_from_table(path.name, header_cells, table_rows, file_warnings, sites, roles)


def read_file(
    path: Path, sites: list[str] | None = None, roles: list[str] | None = None
) -> SheetResult:
    """Le a planilha.

    `sites` e `roles` (opcionais) validam as colunas Site e Papel contra o NetBox.
    """
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_csv(path, sites, roles)
    if suffix == ".xlsx":
        return read_xlsx(path, sites, roles)
    raise ValueError(
        "Formato nao suportado. Use .xlsx ou .csv (se tiver um .xls antigo, salve como .xlsx)."
    )


def _write_list_sheet(workbook, sites: list[str], roles: list[str]) -> None:
    """Aba de apoio com os sites (coluna A) e os papeis (coluna B).

    Fica **visivel e com filtro** de proposito: a lista suspensa da celula serve para
    escolher rapido (digitando, ela pula ao item), mas com milhares de sites achar um
    especifico e mais facil filtrando/buscando aqui - o filtro do Excel tem busca.
    """
    sheet = workbook.create_sheet(LIST_SHEET_TITLE)
    sheet.append(["Sites", "Papeis"])
    for indice in range(max(len(sites), len(roles))):
        sheet.append(
            [
                sites[indice] if indice < len(sites) else None,
                roles[indice] if indice < len(roles) else None,
            ]
        )
    sheet.column_dimensions["A"].width = max((len(item) for item in sites), default=10) + 2
    sheet.column_dimensions["B"].width = max((len(item) for item in roles), default=10) + 2
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:B{max(len(sites), len(roles)) + 1}"


def _add_dropdown(
    sheet,
    selected: list[str],
    key: str,
    list_column: str,
    valores: list[str],
    titulo_erro: str,
    mensagem_erro: str,
) -> None:
    """Lista suspensa numa coluna da planilha, apontando para a aba de listas.

    Nao mexemos em `showDropDown`: no formato do Excel esse atributo e invertido
    (1 = esconder a seta).
    """
    if key not in selected or not valores:
        return
    letra = get_column_letter(selected.index(key) + 1)
    validacao = DataValidation(
        type="list",
        formula1=f"={LIST_SHEET_TITLE}!${list_column}$2:${list_column}${len(valores) + 1}",
        allow_blank=True,
    )
    validacao.errorTitle = titulo_erro
    validacao.error = mensagem_erro
    validacao.showErrorMessage = True
    sheet.add_data_validation(validacao)
    validacao.add(f"{letra}2:{letra}{DROPDOWN_ROWS}")


def write_example_template(
    path: Path,
    sites: list[str] | None = None,
    roles: list[str] | None = None,
    columns: list[str] | None = None,
) -> None:
    """Gera o xlsx de exemplo, com listas suspensas de Site e de Papel.

    Site e Papel sao referencias do NetBox: em vez de digitar, a pessoa escolhe na
    celula - e o app ainda revalida ao carregar.
    """
    selected = template_columns(columns)
    specs = [COLUMNS_BY_KEY[key] for key in selected]

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = DATA_SHEET_TITLE
    sheet.append([spec.label for spec in specs])

    known_sites = sorted({site for site in (sites or []) if site})
    known_roles = sorted({role for role in (roles or []) if role})
    # A amostra usa valores reais (e "Camera" como papel, quando existir).
    papel_amostra = next(
        (role for role in known_roles if normalize_for_match(role) == "camera"),
        known_roles[0] if known_roles else "",
    )
    sample = []
    for spec in specs:
        if spec.key == SITE_KEY and known_sites:
            sample.append(known_sites[0])
        elif spec.key == ROLE_KEY and papel_amostra:
            sample.append(papel_amostra)
        else:
            sample.append(spec.sample)
    sheet.append(sample)

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill

    for column_cells in sheet.columns:
        width = max(
            len(str(cell.value)) for cell in column_cells if cell.value is not None
        ) + 2
        sheet.column_dimensions[column_cells[0].column_letter].width = min(width, 60)

    if known_sites or known_roles:
        _write_list_sheet(workbook, known_sites, known_roles)
        _add_dropdown(
            sheet,
            selected,
            SITE_KEY,
            "A",
            known_sites,
            "Site invalido",
            "Escolha um site que exista no NetBox (use a seta da celula).",
        )
        _add_dropdown(
            sheet,
            selected,
            ROLE_KEY,
            "B",
            known_roles,
            "Papel invalido",
            "Escolha um papel que exista no NetBox (use a seta da celula).",
        )

    workbook.save(path)
    workbook.close()
