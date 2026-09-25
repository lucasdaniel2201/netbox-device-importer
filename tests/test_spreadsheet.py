"""Testes de leitura/validacao/geracao da planilha de cameras (app/spreadsheet.py).

Cobrem o mapeamento de colunas (incluindo a coluna Site), a validacao linha a
linha contra a lista de sites do NetBox e, principalmente, o **dropdown de Site**
no modelo gerado - que e o que evita cadastrar camera em site inexistente.
"""

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402

from app import spreadsheet as sp  # noqa: E402

SITES = ["AME VILANOVA", "CASA DA ESPERANCA", "PSF 1", "PSF 4"]
ROLES = ["Camera", "Switch", "NVR", "Access Point"]
HEADER = [sp.COLUMNS_BY_KEY[key].label for key in sp.TARGET_FIELDS]


def write_sheet(path: Path, header: list[str], rows: list[list[str]]) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def row(**values) -> list[str]:
    """Linha completa na ordem canonica, preenchendo o resto com vazio."""
    cells = {
        "Name": "CAM-01",
        "Site": "PSF 1",
        "Role": "",
        "IP": "10.0.0.1",
        "Vendor": "Hikvision",
        "Model": "DS-2CD",
        "Firmware": "",
        "MAC address": "",
        "Server": "",
        "Description": "",
    }
    cells.update(values)
    return [cells[key] for key in sp.TARGET_FIELDS]


class TestNormalizacao(unittest.TestCase):
    def test_chave_ignora_acento_e_caixa(self):
        self.assertEqual(sp.normalize_for_match("Casa da Esperança"), "casa da esperanca")

    def test_chave_colapsa_espacos(self):
        self.assertEqual(sp.normalize_for_match("  PSF   1  "), "psf 1")

    def test_vendor_padronizado(self):
        self.assertEqual(sp.normalize_vendor("HIKVISION"), "Hikvision")
        self.assertEqual(sp.normalize_vendor("dahua"), "Dahua")
        self.assertEqual(sp.normalize_vendor("Intelbras"), "Intelbras")

    def test_nome_do_device_nao_perde_acento(self):
        """NetBox aceita UTF-8: o nome do device nao perde acento."""
        self.assertEqual(sp.normalize_device_name("Câmera São João"), "Câmera São João")

    def test_nome_do_device_colapsa_espacos(self):
        self.assertEqual(sp.normalize_device_name("  CAM   01 "), "CAM 01")

    def test_site_index_usa_chave_normalizada(self):
        indice = sp.site_index(SITES)
        self.assertEqual(indice["psf 1"], "PSF 1")
        self.assertEqual(indice["casa da esperanca"], "CASA DA ESPERANCA")


class TestColumnSpecsIntegridade(unittest.TestCase):
    def test_chaves_unicas(self):
        keys = [spec.key for spec in sp.COLUMN_SPECS]
        self.assertEqual(len(keys), len(set(keys)))

    def test_rotulos_unicos(self):
        labels = [spec.label for spec in sp.COLUMN_SPECS]
        self.assertEqual(len(labels), len(set(labels)))

    def test_fixas_sao_name_site_ip(self):
        self.assertEqual(sp.FIXED_COLUMNS, ["Name", "Site", "IP"])

    def test_somente_name_e_site_exigem_valor(self):
        """IP pode vir em branco: a camera fica documentada sem endereco."""
        exigem = [key for key in sp.FIXED_COLUMNS if sp.COLUMNS_BY_KEY[key].required_value]
        self.assertEqual(exigem, ["Name", "Site"])

    def test_apelidos_nao_colidem_entre_colunas(self):
        visto: dict[str, str] = {}
        for spec in sp.COLUMN_SPECS:
            for alias in (spec.label, spec.key, *spec.aliases):
                normalized = sp._normalize_header(alias)
                outro = visto.get(normalized)
                if outro is not None and outro != spec.key:
                    self.fail(f"apelido '{alias}' colide entre '{outro}' e '{spec.key}'")
                visto[normalized] = spec.key

    def test_tem_firmware_e_nvr(self):
        """Firmware e o NVR vao para os comments do device - precisam existir."""
        self.assertIn("Firmware", sp.TARGET_FIELDS)
        self.assertIn("Server", sp.TARGET_FIELDS)


class TestTemplateColumns(unittest.TestCase):
    def test_padrao_inclui_todas(self):
        self.assertEqual(sp.template_columns(), sp.TARGET_FIELDS)

    def test_subconjunto_mantem_fixas_e_ordem(self):
        columns = sp.template_columns(["Firmware"])
        self.assertIn("Name", columns)
        self.assertIn("Site", columns)
        self.assertIn("Firmware", columns)
        self.assertNotIn("Vendor", columns)
        self.assertEqual(columns, [c for c in sp.TARGET_FIELDS if c in columns])

    def test_nao_aceita_remover_fixa(self):
        self.assertEqual(sp.template_columns([]), sp.FIXED_COLUMNS)


class TestWriteExampleTemplate(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def _load(self, path: Path):
        return openpyxl.load_workbook(path)

    def test_cabecalho_usa_rotulos_do_modelo(self):
        path = self.tmp / "modelo.xlsx"
        sp.write_example_template(path)
        header = [c.value for c in self._load(path).active[1]]
        self.assertEqual(header, HEADER)

    def test_linha_de_amostra_usa_site_da_lista(self):
        path = self.tmp / "modelo.xlsx"
        sp.write_example_template(path, sites=SITES)
        sheet = self._load(path).active
        site_column = sp.TARGET_FIELDS.index("Site")
        self.assertIn(sheet.cell(row=2, column=site_column + 1).value, SITES)

    def test_cria_aba_de_listas_com_os_sites(self):
        path = self.tmp / "modelo.xlsx"
        sp.write_example_template(path, sites=SITES)
        workbook = self._load(path)
        self.assertIn(sp.LIST_SHEET_TITLE, workbook.sheetnames)
        list_sheet = workbook[sp.LIST_SHEET_TITLE]
        valores = [list_sheet.cell(row=i, column=1).value for i in range(2, 2 + len(SITES))]
        self.assertEqual(sorted(valores), sorted(SITES))

    def test_aba_de_listas_fica_visivel_e_filtravel(self):
        """Com milhares de sites, o filtro da aba de apoio e onde da para buscar."""
        path = self.tmp / "modelo.xlsx"
        sp.write_example_template(path, sites=SITES, roles=ROLES)
        aba = self._load(path)[sp.LIST_SHEET_TITLE]
        self.assertEqual(aba.sheet_state, "visible")
        self.assertEqual(aba.auto_filter.ref, f"A1:B{max(len(SITES), len(ROLES)) + 1}")
        self.assertEqual(aba.freeze_panes, "A2")

    def test_aba_de_listas_tem_sites_e_papeis(self):
        path = self.tmp / "modelo.xlsx"
        sp.write_example_template(path, sites=SITES, roles=ROLES)
        aba = self._load(path)[sp.LIST_SHEET_TITLE]
        self.assertEqual(aba.cell(row=1, column=1).value, "Sites")
        self.assertEqual(aba.cell(row=1, column=2).value, "Papeis")
        papeis = [aba.cell(row=i, column=2).value for i in range(2, 2 + len(ROLES))]
        self.assertEqual(sorted(papeis), sorted(ROLES))

    def test_dropdown_de_site_e_de_papel(self):
        path = self.tmp / "modelo.xlsx"
        sp.write_example_template(path, sites=SITES, roles=ROLES)
        formulas = sorted(
            v.formula1 for v in self._load(path).active.data_validations.dataValidation
        )
        self.assertEqual(
            formulas,
            [
                f"={sp.LIST_SHEET_TITLE}!$A$2:$A${len(SITES) + 1}",
                f"={sp.LIST_SHEET_TITLE}!$B$2:$B${len(ROLES) + 1}",
            ],
        )

    def test_dropdown_de_papel_fica_na_coluna_papel(self):
        from openpyxl.utils import get_column_letter

        path = self.tmp / "modelo.xlsx"
        sp.write_example_template(path, sites=SITES, roles=ROLES)
        validacoes = self._load(path).active.data_validations.dataValidation
        papel = next(v for v in validacoes if f"$B${len(ROLES) + 1}" in v.formula1)
        letra = get_column_letter(sp.TARGET_FIELDS.index(sp.ROLE_KEY) + 1)
        self.assertEqual(str(papel.sqref), f"{letra}2:{letra}{sp.DROPDOWN_ROWS}")

    def test_amostra_usa_camera_como_papel(self):
        path = self.tmp / "modelo.xlsx"
        sp.write_example_template(path, sites=SITES, roles=ROLES)
        coluna = sp.TARGET_FIELDS.index(sp.ROLE_KEY) + 1
        self.assertEqual(self._load(path).active.cell(row=2, column=coluna).value, "Camera")

    def test_sem_papeis_nao_cria_dropdown_de_papel(self):
        path = self.tmp / "modelo.xlsx"
        sp.write_example_template(path, sites=SITES)
        self.assertEqual(len(self._load(path).active.data_validations.dataValidation), 1)

    def test_dropdown_aponta_para_a_aba_de_listas(self):
        path = self.tmp / "modelo.xlsx"
        sp.write_example_template(path, sites=SITES)
        validations = self._load(path).active.data_validations.dataValidation
        self.assertEqual(len(validations), 1)
        self.assertEqual(validations[0].type, "list")
        self.assertEqual(
            validations[0].formula1, f"={sp.LIST_SHEET_TITLE}!$A$2:$A${len(SITES) + 1}"
        )

    def test_dropdown_esta_na_coluna_site(self):
        from openpyxl.utils import get_column_letter

        path = self.tmp / "modelo.xlsx"
        sp.write_example_template(path, sites=SITES)
        validations = self._load(path).active.data_validations.dataValidation
        esperada = get_column_letter(sp.TARGET_FIELDS.index("Site") + 1)
        self.assertEqual(str(validations[0].sqref), f"{esperada}2:{esperada}{sp.DROPDOWN_ROWS}")

    def test_sem_lista_de_sites_nao_cria_dropdown(self):
        path = self.tmp / "modelo.xlsx"
        sp.write_example_template(path)
        workbook = self._load(path)
        self.assertNotIn(sp.LIST_SHEET_TITLE, workbook.sheetnames)
        self.assertEqual(len(workbook.active.data_validations.dataValidation), 0)

    def test_modelo_gerado_pode_ser_lido_e_evalido(self):
        path = self.tmp / "modelo.xlsx"
        sp.write_example_template(path, sites=SITES)
        sheet = sp.read_file(path, sites=SITES)
        self.assertEqual(len(sheet.rows), 1)
        self.assertEqual(len(sheet.valid_rows), 1)
        self.assertIn(sheet.rows[0].cells["Site"], SITES)


class TestReadAndValidate(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def _read(self, header, rows, sites=SITES, roles=ROLES):
        path = self.tmp / "dados.xlsx"
        write_sheet(path, header, rows)
        return sp.read_file(path, sites=sites, roles=roles)

    def test_apelidos_aceitos(self):
        sheet = self._read(
            ["Nome", "Unidade", "IP/Nome", "Fabricante:", "MAC"],
            [["CAM-01", "PSF 1", "10.0.0.1", "Hikvision", "AA-BB-CC-DD-EE-FF"]],
        )
        cells = sheet.rows[0].cells
        self.assertEqual(cells["Name"], "CAM-01")
        self.assertEqual(cells["Site"], "PSF 1")
        self.assertEqual(cells["IP"], "10.0.0.1")
        self.assertEqual(cells["Vendor"], "Hikvision")
        self.assertEqual(cells["MAC address"], "AA-BB-CC-DD-EE-FF")

    def test_coluna_obrigatoria_ausente(self):
        with self.assertRaises(ValueError) as ctx:
            self._read(["Nome do host", "IP"], [["CAM-01", "10.0.0.1"]])
        self.assertIn("obrigatoria", str(ctx.exception))
        self.assertIn("Site", str(ctx.exception))

    def test_site_inexistente_e_erro(self):
        sheet = self._read(HEADER, [row(Site="SITE QUE NAO EXISTE")])
        self.assertFalse(sheet.rows[0].valid)
        self.assertTrue(any("nao existe no NetBox" in i for i in sheet.rows[0].issues))

    def test_site_com_caixa_e_acento_diferentes_e_aceito(self):
        """A comparacao com o site do NetBox e sem acento e sem caixa."""
        sheet = self._read(HEADER, [row(Site="casa da esperança")])
        self.assertTrue(sheet.rows[0].valid, sheet.rows[0].issues)

    def test_site_vazio_e_erro(self):
        sheet = self._read(HEADER, [row(Site="")])
        self.assertTrue(any("Site vazio" in i for i in sheet.rows[0].issues))

    def test_sem_lista_de_sites_so_exige_preenchido(self):
        sheet = self._read(HEADER, [row(Site="QUALQUER COISA")], sites=None)
        self.assertTrue(sheet.rows[0].valid, sheet.rows[0].issues)

    def test_ip_vazio_e_permitido(self):
        """Camera sem IP ainda assim e documentada (a interface/MAC valem sozinhos)."""
        sheet = self._read(HEADER, [row(IP="")])
        self.assertTrue(sheet.rows[0].valid, sheet.rows[0].issues)

    def test_coluna_do_modelo_ausente_e_erro(self):
        with self.assertRaises(ValueError) as ctx:
            self._read(["Nome do host"], [["CAM-01"]])
        mensagem = str(ctx.exception)
        self.assertIn("obrigatoria", mensagem)
        self.assertIn("Site", mensagem)
        self.assertIn("IP", mensagem)

    def test_ip_invalido_e_erro(self):
        sheet = self._read(HEADER, [row(IP="999.1.1")])
        self.assertTrue(any("formato invalido" in i for i in sheet.rows[0].issues))

    def test_papel_valido_e_aceito(self):
        sheet = self._read(HEADER, [row(Role="Switch")])
        self.assertTrue(sheet.rows[0].valid, sheet.rows[0].issues)
        self.assertEqual(sheet.rows[0].cells["Role"], "Switch")

    def test_papel_inexistente_e_erro(self):
        sheet = self._read(HEADER, [row(Role="Roteador Quântico")])
        self.assertFalse(sheet.rows[0].valid)
        self.assertTrue(any("nao existe no NetBox" in i for i in sheet.rows[0].issues))

    def test_papel_vazio_e_permitido(self):
        """Em branco = usa o papel escolhido na tela (o app resolve depois)."""
        sheet = self._read(HEADER, [row(Role="")])
        self.assertTrue(sheet.rows[0].valid, sheet.rows[0].issues)

    def test_papel_ignora_caixa(self):
        sheet = self._read(HEADER, [row(Role="switch")])
        self.assertTrue(sheet.rows[0].valid, sheet.rows[0].issues)

    def test_sem_lista_de_papeis_so_exige_preenchido(self):
        sheet = self._read(HEADER, [row(Role="Qualquer Coisa")], roles=None)
        self.assertTrue(sheet.rows[0].valid, sheet.rows[0].issues)

    def test_ip_com_octeto_maior_que_255_e_erro(self):
        sheet = self._read(HEADER, [row(IP="10.0.0.300")])
        self.assertTrue(any("formato invalido" in i for i in sheet.rows[0].issues))

    def test_nome_longo_demais_e_erro(self):
        sheet = self._read(HEADER, [row(Name="X" * (sp.MAX_DEVICE_NAME_LENGTH + 1))])
        self.assertTrue(any("caracteres" in i for i in sheet.rows[0].issues))

    def test_nome_no_limite_e_valido(self):
        sheet = self._read(HEADER, [row(Name="X" * sp.MAX_DEVICE_NAME_LENGTH)])
        self.assertTrue(sheet.rows[0].valid, sheet.rows[0].issues)

    def test_nome_duplicado_e_erro_nas_duas_linhas(self):
        sheet = self._read(HEADER, [row(Name="CAM-A"), row(Name="CAM-A", IP="10.0.0.2")])
        self.assertEqual(len(sheet.rows_with_errors), 2)
        self.assertTrue(all("duplicado" in i for r in sheet.rows for i in r.issues))

    def test_duplicado_ignora_caixa(self):
        sheet = self._read(HEADER, [row(Name="CAM-A"), row(Name="cam-a", IP="10.0.0.2")])
        self.assertEqual(len(sheet.rows_with_errors), 2)

    def test_mac_estranho_gera_aviso(self):
        sheet = self._read(HEADER, [row(**{"MAC address": "ZZZZ"})])
        self.assertTrue(sheet.rows[0].valid)
        self.assertTrue(any("MAC" in w for w in sheet.rows[0].warnings))

    def test_mac_valido_nao_gera_aviso(self):
        sheet = self._read(HEADER, [row(**{"MAC address": "AA:BB:CC:DD:EE:FF"})])
        self.assertEqual(sheet.rows[0].warnings, [])

    def test_linha_sem_nome_e_ignorada(self):
        sheet = self._read(HEADER, [row(**{"Name": ""})])
        self.assertEqual(len(sheet.rows), 0)
        self.assertEqual(sheet.skipped_empty, 1)

    def test_vendor_hikvision_padronizado(self):
        sheet = self._read(HEADER, [row(Vendor="HIKVISION")])
        self.assertEqual(sheet.rows[0].cells["Vendor"], "Hikvision")

    def test_csv_tambem_e_lido(self):
        path = self.tmp / "dados.csv"
        path.write_text(
            "Nome do host,Site,IP\nCAM-01,PSF 1,10.0.0.1\n", encoding="utf-8"
        )
        sheet = sp.read_file(path, sites=SITES)
        self.assertEqual(len(sheet.valid_rows), 1)

    def test_csv_avisa_que_nao_tem_lista_suspensa(self):
        path = self.tmp / "dados.csv"
        path.write_text(
            "Nome do host,Site,IP\nCAM-01,PSF 1,10.0.0.1\n", encoding="utf-8"
        )
        sheet = sp.read_file(path, sites=SITES)
        self.assertTrue(sheet.file_warnings)

    def test_formato_nao_suportado(self):
        path = self.tmp / "dados.txt"
        path.write_text("x", encoding="utf-8")
        with self.assertRaises(ValueError):
            sp.read_file(path)


if __name__ == "__main__":
    unittest.main(verbosity=2)
