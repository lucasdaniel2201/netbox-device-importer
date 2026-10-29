"""Testes da planilha de links IP (app/link_sheet.py).

O que precisa valer: o CIDR e normalizado antes de qualquer envio (o NetBox
compara `prefix` como texto), o IP do link tem que pertencer ao bloco publico, e
bloco/IP repetido em duas linhas e erro nas duas - senao a segunda linha
sobrescreveria a primeira em silencio.
"""

import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

import ipaddress  # noqa: E402

import openpyxl  # noqa: E402

from app import link_sheet as ls  # noqa: E402

SITES = ["SITE A", "SITE EXEMPLO"]
ROLES = ["Reservado", "Servidores"]
HEADER = [ls.COLUMNS_BY_KEY[key].label for key in ls.TARGET_FIELDS]


def write_sheet(path: Path, header: list[str], rows: list[list[str]]) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def row(**values) -> list[str]:
    """Linha completa na ordem canonica, preenchendo o resto com vazio."""
    cells = {
        ls.SITE_KEY: "SITE EXEMPLO",
        ls.DEVICE_KEY: "RT-DELEGACIA-GERAL",
        ls.VENDOR_KEY: "MikroTik",
        ls.MODEL_KEY: "RB760iGS",
        ls.DEVICE_ROLE_KEY: "Roteador",
        ls.PUBLIC_KEY: "203.0.113.240/30",
        ls.INTERFACE_KEY: "wan1",
        ls.IP_KEY: "",
        ls.LAN_KEY: "10.10.48.0/24",
        ls.ROLE_KEY: "",
        ls.ZABBIX_KEY: "",
        ls.DESCRIPTION_KEY: "PROVEDOR",
        ls.PROVIDER_KEY: "",
    }
    cells.update(values)
    return [cells[key] for key in ls.TARGET_FIELDS]


class TestCidr(unittest.TestCase):
    def test_normaliza_host_para_a_rede(self):
        """A planilha as vezes traz o endereco do roteador na coluna do bloco."""
        self.assertEqual(
            str(ls.parse_network("203.0.113.241/30", "bloco público")), "203.0.113.240/30"
        )

    def test_cidr_sem_mascara_e_erro(self):
        with self.assertRaises(ValueError):
            ls.parse_network("203.0.113.240", "bloco público")

    def test_cidr_invalido_e_erro(self):
        with self.assertRaises(ValueError):
            ls.parse_network("203.0.113.999/30", "bloco público")

    def test_ipv6_nao_e_aceito(self):
        with self.assertRaises(ValueError) as ctx:
            ls.parse_network("2001:db8::/64", "bloco público")
        self.assertIn("IPv4", str(ctx.exception))

    def test_ip_do_link_fora_do_bloco(self):
        rede = ipaddress.ip_network("203.0.113.240/30")
        with self.assertRaises(ValueError) as ctx:
            ls.parse_host_address("10.0.0.1", rede)
        self.assertIn("nao pertence", str(ctx.exception))

    def test_ip_do_link_aceita_com_e_sem_mascara(self):
        rede = ipaddress.ip_network("203.0.113.240/30")
        self.assertEqual(str(ls.parse_host_address("203.0.113.241", rede)), "203.0.113.241")
        self.assertEqual(str(ls.parse_host_address("203.0.113.241/30", rede)), "203.0.113.241")


class TestEnderecoDeduzido(unittest.TestCase):
    """/30 deduz rede+1; /31 e ponto-a-ponto (RFC 3021) e o cliente e o segundo."""

    def test_barra_30(self):
        self.assertEqual(
            str(ls.default_link_ip(ipaddress.ip_network("203.0.113.240/30"))),
            "203.0.113.241",
        )

    def test_barra_31(self):
        self.assertEqual(
            str(ls.default_link_ip(ipaddress.ip_network("198.51.100.224/31"))),
            "198.51.100.225",
        )

    def test_barra_32(self):
        self.assertEqual(
            str(ls.default_link_ip(ipaddress.ip_network("203.0.113.123/32"))),
            "203.0.113.123",
        )

    def test_rede_e_broadcast_nao_sao_utilizaveis(self):
        rede = ipaddress.ip_network("10.0.0.0/24")
        self.assertFalse(ls.usable_host_address(rede, rede.network_address))
        self.assertFalse(ls.usable_host_address(rede, rede.broadcast_address))
        self.assertTrue(ls.usable_host_address(rede, ipaddress.ip_address("10.0.0.7")))

    def test_barra_31_aceita_os_dois(self):
        rede = ipaddress.ip_network("198.51.100.224/31")
        self.assertTrue(ls.usable_host_address(rede, rede.network_address))
        self.assertTrue(ls.usable_host_address(rede, rede[1]))


class TestColunas(unittest.TestCase):
    def test_chaves_e_rotulos_unicos(self):
        chaves = [spec.key for spec in ls.COLUMN_SPECS]
        rotulos = [spec.label for spec in ls.COLUMN_SPECS]
        self.assertEqual(len(chaves), len(set(chaves)))
        self.assertEqual(len(rotulos), len(set(rotulos)))

    def test_fixas_sao_site_device_bloco(self):
        self.assertEqual(
            ls.FIXED_COLUMNS, [ls.SITE_KEY, ls.DEVICE_KEY, ls.PUBLIC_KEY]
        )

    def test_apelidos_nao_colidem(self):
        visto: dict[str, str] = {}
        for spec in ls.COLUMN_SPECS:
            for alias in (spec.label, spec.key, *spec.aliases):
                normalizado = ls.normalize_for_match(alias)
                outro = visto.get(normalizado)
                if outro is not None and outro != spec.key:
                    self.fail(f"apelido '{alias}' colide entre '{outro}' e '{spec.key}'")
                visto[normalizado] = spec.key


class TestModelo(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def _load(self, path: Path):
        return openpyxl.load_workbook(path)

    def test_cabecalho_usa_os_rotulos(self):
        path = self.tmp / "links.xlsx"
        ls.write_example_template(path)
        self.assertEqual([c.value for c in self._load(path).active[1]], HEADER)

    def test_listas_de_site_e_papel(self):
        path = self.tmp / "links.xlsx"
        ls.write_example_template(path, sites=SITES, prefix_roles=ROLES)
        aba = self._load(path)[ls.LIST_SHEET_TITLE]
        self.assertEqual(aba.cell(row=1, column=1).value, "Sites")
        self.assertEqual(aba.cell(row=1, column=2).value, "Papeis de prefixo")
        self.assertEqual(aba.cell(row=2, column=1).value, "SITE A")

    def test_dropdowns_de_site_papel_e_device(self):
        path = self.tmp / "links.xlsx"
        ls.write_example_template(
            path, sites=SITES, prefix_roles=ROLES, devices=["RT-1", "RT-2"]
        )
        formulas = sorted(
            v.formula1 for v in self._load(path).active.data_validations.dataValidation
        )
        self.assertEqual(
            formulas,
            [
                f"={ls.LIST_SHEET_TITLE}!$A$2:$A${len(SITES) + 1}",
                f"={ls.LIST_SHEET_TITLE}!$B$2:$B${len(ROLES) + 1}",
                f"={ls.LIST_SHEET_TITLE}!$C$2:$C$3",
            ],
        )

    def test_dropdown_de_device_fica_na_coluna_device(self):
        from openpyxl.utils import get_column_letter

        path = self.tmp / "links.xlsx"
        ls.write_example_template(
            path, sites=SITES, prefix_roles=ROLES, devices=["RT-1", "RT-2"]
        )
        validacoes = self._load(path).active.data_validations.dataValidation
        device = next(v for v in validacoes if "!$C$" in v.formula1)
        letra = get_column_letter(ls.TARGET_FIELDS.index(ls.DEVICE_KEY) + 1)
        self.assertEqual(str(device.sqref), f"{letra}2:{letra}{ls.DROPDOWN_ROWS}")

    def test_aba_de_listas_tem_os_devices(self):
        path = self.tmp / "links.xlsx"
        ls.write_example_template(
            path, sites=SITES, prefix_roles=ROLES, devices=["RT-1", "RT-2"]
        )
        aba = self._load(path)[ls.LIST_SHEET_TITLE]
        self.assertEqual(aba.cell(row=1, column=3).value, "Devices")
        self.assertEqual(aba.cell(row=2, column=3).value, "RT-1")

    def test_amostra_usa_o_site_real(self):
        path = self.tmp / "links.xlsx"
        ls.write_example_template(path, sites=SITES, prefix_roles=ROLES)
        valores = [c.value for c in self._load(path).active[2]]
        linha = dict(zip(ls.TARGET_FIELDS, valores, strict=True))
        self.assertEqual(linha[ls.SITE_KEY], "SITE A")

    def test_papel_da_amostra_so_quando_e_papel_de_link(self):
        """Herdar o primeiro papel ("Reservado") marcaria o bloco com o papel errado."""
        path = self.tmp / "links.xlsx"
        ls.write_example_template(path, sites=SITES, prefix_roles=["Reservado", "Servidores"])
        valores = [c.value for c in self._load(path).active[2]]
        linha = dict(zip(ls.TARGET_FIELDS, valores, strict=True))
        self.assertFalse(linha[ls.ROLE_KEY])

        outro = self.tmp / "links-com-papel.xlsx"
        ls.write_example_template(outro, sites=SITES, prefix_roles=["Link WAN"])
        valores = [c.value for c in self._load(outro).active[2]]
        linha = dict(zip(ls.TARGET_FIELDS, valores, strict=True))
        self.assertEqual(linha[ls.ROLE_KEY], "Link WAN")

    def test_modelo_gerado_pode_ser_lido(self):
        path = self.tmp / "links.xlsx"
        ls.write_example_template(path, sites=SITES, prefix_roles=ROLES)
        sheet = ls.read_file(path, sites=SITES, prefix_roles=ROLES)
        self.assertEqual(len(sheet.rows), 1)
        self.assertTrue(sheet.rows[0].valid, sheet.rows[0].issues)
        self.assertEqual(sheet.rows[0].cells[ls.PUBLIC_KEY], "203.0.113.240/30")


class TestLeituraEValidacao(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.tmp = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def _read(self, header, rows, sites=SITES, prefix_roles=ROLES):
        path = self.tmp / "dados.xlsx"
        write_sheet(path, header, rows)
        return ls.read_file(path, sites=sites, prefix_roles=prefix_roles)

    def test_linha_valida(self):
        sheet = self._read(HEADER, [row()])
        linha = sheet.rows[0]
        self.assertTrue(linha.valid, linha.issues)
        self.assertEqual(linha.cells[ls.PUBLIC_KEY], "203.0.113.240/30")

    def test_apelidos_aceitos(self):
        """A planilha real usa outros nomes: Delegacia=Site e Equipamento=Modelo."""
        sheet = self._read(
            ["Delegacia", "Dispositivo", "Equipamento", "Bloco de IP público", "Bloco lan", "Des. parceiro"],
            [["SITE EXEMPLO", "RT-1", "RB760iGS", "203.0.113.240/30", "10.10.48.0/24", "PROVEDOR"]],
        )
        linha = sheet.rows[0]
        self.assertEqual(linha.cells[ls.SITE_KEY], "SITE EXEMPLO")
        self.assertEqual(linha.cells[ls.DEVICE_KEY], "RT-1")
        self.assertEqual(linha.cells[ls.MODEL_KEY], "RB760iGS")
        self.assertEqual(linha.cells[ls.LAN_KEY], "10.10.48.0/24")
        self.assertEqual(linha.cells[ls.PROVIDER_KEY], "PROVEDOR")
        self.assertTrue(linha.valid, linha.issues)

    def test_coluna_obrigatoria_ausente(self):
        with self.assertRaises(ValueError) as ctx:
            self._read(["Site", "Device"], [["SITE EXEMPLO", "RT-1"]])
        self.assertIn("Bloco público", str(ctx.exception))

    def test_site_inexistente(self):
        sheet = self._read(HEADER, [row(**{ls.SITE_KEY: "SITE FANTASMA"})])
        self.assertFalse(sheet.rows[0].valid)
        self.assertTrue(any("nao existe no NetBox" in i for i in sheet.rows[0].issues))

    def test_site_ignora_caixa(self):
        """A comparacao com o site do NetBox e sem acento e sem caixa."""
        sheet = self._read(HEADER, [row(**{ls.SITE_KEY: "site a"})])
        self.assertTrue(sheet.rows[0].valid, sheet.rows[0].issues)

    def test_device_vazio_e_erro(self):
        sheet = self._read(HEADER, [row(**{ls.DEVICE_KEY: ""})])
        self.assertTrue(any("Device vazio" in i for i in sheet.rows[0].issues))

    def test_bloco_publico_vazio_e_erro(self):
        sheet = self._read(HEADER, [row(**{ls.PUBLIC_KEY: ""})])
        self.assertTrue(any("Bloco público vazio" in i for i in sheet.rows[0].issues))

    def test_bloco_invalido_e_erro(self):
        sheet = self._read(HEADER, [row(**{ls.PUBLIC_KEY: "203.0.113.999/30"})])
        self.assertTrue(any("CIDR invalido" in i for i in sheet.rows[0].issues))

    def test_bloco_com_endereco_do_roteador_e_ajustado(self):
        """A planilha real traz `203.0.113.241/30`: o bloco e .240/30 e o .241 e o roteador."""
        sheet = self._read(HEADER, [row(**{ls.PUBLIC_KEY: "203.0.113.241/30"})])
        linha = sheet.rows[0]
        self.assertTrue(linha.valid, linha.issues)
        self.assertEqual(linha.cells[ls.PUBLIC_KEY], "203.0.113.240/30")
        self.assertEqual(linha.block_host, "203.0.113.241")
        self.assertTrue(any("endereço do roteador" in w for w in linha.warnings))

    def test_bloco_em_slash_31_com_endereco_do_roteador(self):
        """Caso real: 198.51.100.45/31 -> rede .44/31, roteador .45."""
        sheet = self._read(HEADER, [row(**{ls.PUBLIC_KEY: "198.51.100.45/31"})])
        linha = sheet.rows[0]
        self.assertTrue(linha.valid, linha.issues)
        self.assertEqual(linha.cells[ls.PUBLIC_KEY], "198.51.100.44/31")
        self.assertEqual(linha.block_host, "198.51.100.45")

    def test_bloco_ja_na_rede_nao_tem_endereco_de_roteador(self):
        sheet = self._read(HEADER, [row()])
        self.assertEqual(sheet.rows[0].block_host, "")

    def test_bloco_lan_repetido_e_aviso_nao_erro(self):
        """Dois links do mesmo site compartilham o /24 (acontece na planilha real)."""
        sheet = self._read(
            HEADER,
            [
                row(**{ls.LAN_KEY: "10.10.48.0/24"}),
                row(
                    **{
                        ls.PUBLIC_KEY: "203.0.113.244/30",
                        ls.IP_KEY: "203.0.113.245",
                        ls.LAN_KEY: "10.10.48.0/24",
                    }
                ),
            ],
        )
        self.assertEqual(sheet.rows_with_errors, [])
        for linha in sheet.rows:
            self.assertTrue(any("Bloco LAN repetido" in w for w in linha.warnings))

    def test_ip_do_link_fora_do_bloco_e_erro(self):
        sheet = self._read(HEADER, [row(**{ls.IP_KEY: "10.0.0.1"})])
        self.assertTrue(any("nao pertence" in i for i in sheet.rows[0].issues))

    def test_ip_de_rede_do_bloco_e_erro(self):
        sheet = self._read(HEADER, [row(**{ls.PUBLIC_KEY: "10.0.0.0/24", ls.IP_KEY: "10.0.0.0"})])
        self.assertTrue(any("rede/broadcast" in i for i in sheet.rows[0].issues))

    def test_ip_sem_mascara_recebe_a_do_bloco(self):
        sheet = self._read(HEADER, [row(**{ls.IP_KEY: "203.0.113.242"})])
        linha = sheet.rows[0]
        self.assertTrue(linha.valid, linha.issues)
        self.assertEqual(linha.cells[ls.IP_KEY], "203.0.113.242/30")
        self.assertTrue(any("sem mascara" in w for w in linha.warnings))

    def test_papel_inexistente_e_erro(self):
        sheet = self._read(HEADER, [row(**{ls.ROLE_KEY: "Papel Quântico"})])
        self.assertTrue(any("nao existe no IPAM" in i for i in sheet.rows[0].issues))

    def test_papel_vazio_e_aceito(self):
        sheet = self._read(HEADER, [row(**{ls.ROLE_KEY: ""})])
        self.assertTrue(sheet.rows[0].valid, sheet.rows[0].issues)

    def test_zabbix_aceita_zero_e_um(self):
        for valor in ("0", "1"):
            with self.subTest(valor=valor):
                sheet = self._read(HEADER, [row(**{ls.ZABBIX_KEY: valor})])
                self.assertTrue(sheet.rows[0].valid, sheet.rows[0].issues)

    def test_zabbix_invalido_e_erro(self):
        sheet = self._read(HEADER, [row(**{ls.ZABBIX_KEY: "sim"})])
        self.assertTrue(any("Zabbix" in i for i in sheet.rows[0].issues))

    def test_bloco_repetido_e_erro_nas_duas_linhas(self):
        sheet = self._read(
            HEADER,
            [row(**{ls.IP_KEY: "203.0.113.241"}), row(**{ls.IP_KEY: "203.0.113.242"})],
        )
        self.assertEqual(len(sheet.rows_with_errors), 2)
        for linha in sheet.rows:
            with self.subTest(linha=linha.row_number):
                self.assertTrue(any("Bloco público repetido" in i for i in linha.issues))

    def test_bloco_repetido_com_mascara_diferente_tambem_conta(self):
        """241/30 e 240/30 sao o mesmo bloco: a duplicidade pega depois de normalizar."""
        sheet = self._read(
            HEADER,
            [row(), row(**{ls.PUBLIC_KEY: "203.0.113.241/30", ls.LAN_KEY: "10.10.49.0/24"})],
        )
        self.assertEqual(len(sheet.rows_with_errors), 2)

    def test_ip_repetido_e_erro(self):
        """Mesmo IP em duas linhas: o segundo sobrescreveria o primeiro em silencio."""
        sheet = self._read(
            HEADER,
            [
                row(**{ls.IP_KEY: "203.0.113.241", ls.LAN_KEY: "10.10.48.0/24"}),
                row(**{ls.IP_KEY: "203.0.113.241/30", ls.LAN_KEY: "10.10.49.0/24"}),
            ],
        )
        self.assertTrue(
            any("IP do link repetido" in i for r in sheet.rows for i in r.issues)
        )

    def test_linha_vazia_e_ignorada(self):
        sheet = self._read(
            HEADER,
            [
                [
                    "",
                ]
                * len(ls.TARGET_FIELDS),
                row(),
            ],
        )
        self.assertEqual(len(sheet.rows), 1)
        self.assertEqual(sheet.skipped_empty, 1)

    def test_csv_tambem_e_lido(self):
        path = self.tmp / "links.csv"
        path.write_text(
            "Site,Device,Bloco público\nSITE EXEMPLO,RT-1,203.0.113.240/30\n",
            encoding="utf-8",
        )
        sheet = ls.read_file(path, sites=SITES)
        self.assertEqual(len(sheet.valid_rows), 1)

    def test_formato_nao_suportado(self):
        path = self.tmp / "links.txt"
        path.write_text("x", encoding="utf-8")
        with self.assertRaises(ValueError):
            ls.read_file(path)


if __name__ == "__main__":
    unittest.main(verbosity=2)
